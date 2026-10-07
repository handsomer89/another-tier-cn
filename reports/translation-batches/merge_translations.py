"""Merge evidence-backed CN-name batches into the audit workbook and CSV files."""
from pathlib import Path
from collections import Counter
import csv
import json
import re
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter

BASE = Path(__file__).resolve().parent
OUT = BASE.parent
ROOT = OUT.parent

def read_csv(name):
    with (BASE/name).open(encoding='utf-8-sig', newline='') as f:
        return list(csv.DictReader(f))

main = read_csv('original-main.csv')
extra = read_csv('original-extra.csv')
all_entries = {i:r for i,r in enumerate(main+extra,1)}
eligible = set()
for i,r in all_entries.items():
    if r['类别']!='角色名' and any(re.search('[\u4e00-\u9fff]',c) for c in r['关联角色/形态'].split('；')):
        eligible.add(i)

latest = {}
history = []
batch_stats = []
for path in sorted(BASE.glob('result-[0-9][0-9].json')):
    batch = path.stem.split('-')[-1]
    input_path = BASE/f'batch-{batch}.json'
    inputs = json.loads(input_path.read_text())
    outputs = json.loads(path.read_text())
    assert len(inputs)==20 and len(outputs)==20, (path,len(inputs),len(outputs))
    expected = {r['id']:r for r in inputs}
    assert len(expected)==20
    assert {r['id'] for r in outputs}==set(expected), path
    assert len({r['id'] for r in outputs})==20
    for r in outputs:
        i = r['id']
        assert i in eligible, (path,i,'excluded item translated')
        assert r['english']==all_entries[i]['英文原文'], (path,i)
        if r.get('chinese'):
            assert isinstance(r['chinese'],str) and r['chinese']==r['chinese'].strip()
            assert r.get('sources'), (path,i,'missing source')
            assert all(s.get('url','').startswith('https://') and s.get('evidence') for s in r['sources']), (path,i)
        latest[i] = dict(r,batch=batch)
        history.append(dict(r,batch=batch))
    batch_stats.append([batch,len(inputs),sum(bool(r.get('chinese')) for r in outputs),sum(not r.get('chinese') for r in outputs),input_path.name,path.name])

root_review = BASE/'root-review.json'
if root_review.exists():
    for r in json.loads(root_review.read_text()):
        assert r['id'] in eligible and r['english']==all_entries[r['id']]['英文原文']
        latest[r['id']]=dict(r,batch='主代理复核')
        history.append(dict(r,batch='主代理复核'))

# All eligible entries must be reviewed, including entries without a verified name.
assert set(latest)==eligible, (eligible-set(latest),set(latest)-eligible)
headers=['条目ID','类别','英文原文','中文译名（国服原文）','核对状态','国服关联角色（按网站）','候选译名/差异说明','来源链接','核对证据','核对批次','出现页数','关联角色/形态','掉落地点或关联材料','难度（页面已汉化）','页面字段','源文件（相对项目根目录）','原始备注']
def format_row(i,r):
    cnchars=[c for c in r['关联角色/形态'].split('；') if re.search('[\u4e00-\u9fff]',c)]
    if i not in eligible:
        status='保留英文（按网站角色名规则）'
        translated='';note='关联角色在网站中均为英文标题；按用户规则不翻译。'
        sources=[];batch='未分配'
    else:
        result=latest[i]
        translated=result.get('chinese','')
        status='国服资料已核对' if translated else ('来源冲突待核实' if '冲突' in result.get('status','') or '冲突' in result.get('note','') else '待核实国服原文')
        note=result.get('note','')
        candidates=result.get('candidates',[])
        if candidates:
            note='候选：'+json.dumps(candidates,ensure_ascii=False)+'\n'+note
        sources=result.get('sources',[])
        batch=result['batch']
    links='\n'.join(dict.fromkeys(s['url'] for s in sources))
    evidence='\n'.join(f"{s.get('title','')}：{s.get('evidence','')}" for s in sources)
    return [i,r['类别'],r['英文原文'],translated,status,'；'.join(cnchars),note,links,evidence,batch,r['出现页数'],r['关联角色/形态'],r['掉落地点或关联材料'],r['难度（页面已汉化）'],r['页面字段'],r['源文件（相对项目根目录）'],r['备注']]

main_rows=[format_row(i,r) for i,r in enumerate(main,1)]
extra_rows=[format_row(i,r) for i,r in enumerate(extra,307)]
pending=[row for row in main_rows+extra_rows if row[0] in eligible and not row[3]]
excluded=[row for row in main_rows+extra_rows if row[0] not in eligible]
count_main=Counter(row[4] for row in main_rows)
count_extra=Counter(row[4] for row in extra_rows)
wb=Workbook();intro=wb.active;intro.title='检查说明'
intro_rows=[
    ['项目','another-tier-cn'],['检查日期','2026-10-07'],
    ['用户指定范围','仅翻译网站中角色标题为中文的关联条目；角色标题纯英文者保留英文。共享副本/材料只要关联至少一个中文标题角色即可纳入。此规则不另行推断国服实装进度。'],
    ['翻译原则','国服中文名称采用公开国服公告、国服游戏截图、国服玩家实际掉落记录、国服BWIKI及明确国服材料表；禁止自行直译、台服转换及自定后缀。'],
    ['主表范围','306项去重职业书、异节、改典、典录、诗篇、史籍及副本/掉落地点，关联322个静态角色详情页。'],
    ['主表已填写',sum(bool(row[3]) for row in main_rows)],
    ['主表待核实',sum(row[0] in eligible and not row[3] for row in main_rows)],
    ['主表按规则保留英文',sum(row[0] not in eligible for row in main_rows)],
    ['附加表已填写',sum(bool(row[3]) for row in extra_rows)],
    ['附加表待核实',sum(row[0] in eligible and not row[3] for row in extra_rows)],
    ['附加表按规则保留英文',sum(row[0] not in eligible for row in extra_rows)],
    ['协作方式','6个子代理，每次20条；前15批为初轮，16批为附加项及补缺，后续为20条复核批。所有分批输入/输出保存在translation-batches。'],
    ['状态含义','国服资料已核对：已找到对应名称原文及来源，并经整合审查；不等于每条都有官方材料截图。待核实/冲突：中文主列留空，已知候选或差异写入说明，不当作确定译名。'],
    ['来源优先与消歧','国服游戏截图及官方材料原文优先；官方职阶公告仅用于交叉核对词根，不能单独确定异节/改典/典录后缀。旧攻略有笔误、BWIKI部分道具来自通用模板，须核实际材料。'],
    ['保留原始信息','英文原文、关联角色、掉落地点、难度及源文件均保留；英文关联角色在中文译名范围之外。未读线上API或修改网站代码。'],
    ['界面枚举说明','Paid是网站自定义枚举，不能当作游戏国服道具名直译；在附加表中单独留待核实。'],
]
for row in intro_rows:intro.append(row)
intro.column_dimensions['A'].width=28;intro.column_dimensions['B'].width=115
for row in intro:
    for c in row:c.alignment=Alignment(wrap_text=True,vertical='top')
    intro.row_dimensions[row[0].row].height=48

def sheet(title,columns,data,table_name,widths=None):
    ws=wb.create_sheet(title);ws.append(columns)
    for row in data:ws.append(row)
    ws.freeze_panes='D2'
    if data:
        table=Table(displayName=table_name,ref=f'A1:{get_column_letter(len(columns))}{len(data)+1}')
        table.tableStyleInfo=TableStyleInfo(name='TableStyleMedium2',showRowStripes=True)
        ws.add_table(table)
    for idx,w in enumerate(widths or [20]*len(columns),1):ws.column_dimensions[get_column_letter(idx)].width=w
    for row in ws:
        for c in row:c.alignment=Alignment(wrap_text=True,vertical='top')
        ws.row_dimensions[row[0].row].height=48 if row[0].row>1 else 32
    for c in ws[1]:c.fill=PatternFill('solid',fgColor='245A81');c.font=Font(bold=True,color='FFFFFF')
    if columns==headers:
        for row in ws.iter_rows(min_row=2):
            row[3].fill=PatternFill('solid',fgColor='E2F0D9' if row[3].value else 'FFF2CC')
    return ws
widths=[10,23,45,32,30,48,65,70,85,15,12,50,60,25,22,65,50]
sheet('国服中文译名主表',headers,main_rows,'MainCNNames',widths)
sheet('其他英文名称核对',headers,extra_rows,'OtherCNNames',widths)
sheet('待核实条目',headers,pending,'PendingCNNames',widths)
sheet('按规则保留英文',headers,excluded,'ExcludedCNNames',widths)
evidence=[]
for r in history:
    for s in r.get('sources',[]):
        evidence.append([r['id'],r['english'],r.get('chinese',''),r['batch'],s.get('title',''),s.get('url',''),s.get('evidence','')])
sheet('来源证据明细',['条目ID','英文原文','该批译名','批次','来源标题','来源URL','证据摘录'],evidence,'SourceEvidence',[10,45,32,15,60,85,100])
sheet('每批20条记录',['批次','条目数','该批填写数','该批待核实数','输入文件','输出文件'],batch_stats,'BatchLog',[12,12,18,18,30,30])
old=load_workbook(BASE/'original-workbook.xlsx',read_only=True)
old_rows=list(old['已有映射保留英文'].values)
sheet('已有映射保留英文',list(old_rows[0]),[list(r) for r in old_rows[1:]],'OriginalPreservedEnglish',[23,45,28,12,50,60,25,22,65,55])
old.close()
wb.save(OUT/'未汉化名称清单.xlsx')
for filename,data in [('未汉化名称清单.csv',main_rows),('其他英文名称待核对.csv',extra_rows),('待核实国服译名.csv',pending)]:
    with (OUT/filename).open('w',encoding='utf-8-sig',newline='') as f:
        writer=csv.writer(f);writer.writerow(headers);writer.writerows(data)
def esc(s):return str(s or '').replace('|','\\|').replace('\n','<br>')
md=['# 当前网站国服中文译名核对表','',
    '检查日期：2026-10-07。仅翻译网站中关联角色标题含中文的条目，纯英文角色相关条目按用户规则保留。',
    f'主表306项：{sum(bool(r[3]) for r in main_rows)}项已核对并填入国服译名，{sum(r[0] in eligible and not r[3] for r in main_rows)}项待核实，{sum(r[0] not in eligible for r in main_rows)}项按规则保留英文。',
    '不自行直译或统一“的／之”后缀。待核实项的候选和来源差异见Excel；原始英文清单保存在translation-batches。',
    'Excel含关联角色、掉落地点、全部来源证据、分批记录和待核实清单。未修改网站代码。','',
    '| ID | 类别 | 英文原文 | 国服中文原文 | 状态 | 首要来源 |', '| ---: | --- | --- | --- | --- | --- |']
for row in main_rows:
    url=row[7].splitlines()[0] if row[7] else ''
    md.append('| '+' | '.join([str(row[0]),esc(row[1]),esc(row[2]),esc(row[3]),esc(row[4]),f'[来源]({url})' if url else ''])+' |')
(OUT/'未汉化名称清单.md').write_text('\n'.join(md)+'\n')
summary={'main':dict(count_main),'extra':dict(count_extra),'main_filled':sum(bool(r[3]) for r in main_rows),'main_pending':sum(r[0] in eligible and not r[3] for r in main_rows),'main_excluded':sum(r[0] not in eligible for r in main_rows),'extra_filled':sum(bool(r[3]) for r in extra_rows),'eligible_unique':len(eligible),'completed_batches':len(batch_stats),'assignments':sum(s[1] for s in batch_stats),'evidence_rows':len(evidence)}
(BASE/'merge-summary.json').write_text(json.dumps(summary,ensure_ascii=False,indent=2))
print(json.dumps(summary,ensure_ascii=False,indent=2))
